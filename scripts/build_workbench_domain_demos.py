#!/usr/bin/env python3
"""Original synthetic domain fixtures; expectations are authored mutation labels.

No checker or LLM is called by this builder. Policy bounds are demonstration
contracts, not clinical, regulatory, financial or manufacturing safety guidance.
"""
from __future__ import annotations

import argparse
import copy
import csv
from datetime import datetime
import io
import json
import re
from pathlib import Path
import sqlite3
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from rule_workbench.core import new_project, validate_project
from rule_workbench.exports import export_json


SPECS = {
    'manufacturing': {
        'primary': 'WorkOrders', 'id': 'work_order_id', 'number': 'planned_units',
        'bounds': [1, 1000], 'enum': 'production_status', 'values': ['planned', 'released', 'closed'],
        'dates': ['scheduled_start', 'scheduled_end'], 'optional': 'setup_minutes',
        'refs': [('Products', 'product_id'), ('Routing', 'routing_id')],
        'conditional': 'If production_status is released, the selected routing must be approved for the selected product.',
        'url': 'https://learn.microsoft.com/en-us/sql/samples/adventureworks-install-configure?view=sql-server-ver17',
    },
    'healthcare': {
        'primary': 'Encounters', 'id': 'encounter_id', 'number': 'administrative_sequence',
        'bounds': [1, 1000], 'enum': 'encounter_status', 'values': ['planned', 'in-progress', 'finished'],
        'dates': ['period_start', 'period_end'], 'optional': 'document_count',
        'refs': [('Patients', 'patient_id'), ('Practitioners', 'practitioner_id')],
        'conditional': 'If encounter_status is finished, a discharge_summary_id must be recorded.',
        'url': 'https://hl7.org/fhir/R4/encounter.html',
    },
    'finance': {
        'primary': 'Payments', 'id': 'payment_id', 'number': 'amount_minor_units',
        'bounds': [1, 100000], 'enum': 'payment_status', 'values': ['pending', 'settled', 'cancelled'],
        'dates': ['initiated_date', 'value_date'], 'optional': 'fee_minor_units',
        'refs': [('Accounts', 'account_id'), ('Counterparties', 'counterparty_id')],
        'conditional': 'If payment_status is settled, settlement_reference must be recorded.',
        'url': 'https://docs.stripe.com/api/payment_intents/object',
    },
    'ecommerce': {
        'primary': 'Orders', 'id': 'order_id', 'number': 'quantity',
        'bounds': [1, 100], 'enum': 'order_status', 'values': ['placed', 'shipped', 'cancelled'],
        'dates': ['order_date', 'promised_ship_date'], 'optional': 'discount_percent',
        'refs': [('Customers', 'customer_id'), ('Products', 'product_id')],
        'conditional': 'If order_status is shipped, tracking_reference must be recorded.',
        'url': 'https://github.com/microsoft/sql-server-samples/blob/master/samples/databases/wide-world-importers/README.md',
    },
}


def _references(domain, spec):
    """Domain attributes and linked records, not just renamed scalar columns."""
    if domain == 'manufacturing':
        return [
            [{'product_id': f'P{i:03}', 'product_name': f'Pump housing {i:02}', 'material_grade': ['AL6061', 'SS304'][i % 2], 'drawing_revision': f'R{i % 3 + 1}'} for i in range(1, 25)],
            [{'routing_id': f'R{i:03}', 'product_id': f'P{i:03}', 'operation_sequence': 'cut>mill>inspect', 'work_center': f'WC{i % 4 + 1}', 'approval_status': 'approved'} for i in range(1, 25)],
        ]
    if domain == 'healthcare':
        return [
            [{'patient_id': f'PT{i:03}', 'synthetic_alias': f'Synthetic patient {i:02}', 'administrative_region': f'Area-{i % 4 + 1}'} for i in range(1, 41)],
            [{'practitioner_id': f'PR{i:03}', 'synthetic_alias': f'Synthetic practitioner {i:02}', 'service_unit': ['outpatient', 'observation'][i % 2]} for i in range(1, 13)],
        ]
    if domain == 'finance':
        return [
            [{'account_id': f'AC{i:03}', 'ledger_name': f'Synthetic operating ledger {i:02}', 'currency': 'USD', 'account_state': 'open'} for i in range(1, 21)],
            [{'counterparty_id': f'CP{i:03}', 'synthetic_name': f'Synthetic trade partner {i:02}', 'relationship_type': ['supplier', 'customer'][i % 2]} for i in range(1, 31)],
        ]
    return [
        [{'customer_id': f'CU{i:03}', 'synthetic_name': f'Synthetic customer {i:02}', 'sales_channel': ['web', 'mobile'][i % 2]} for i in range(1, 41)],
        [{'product_id': f'SKU{i:03}', 'product_name': f'Desk organizer {i:02}', 'category': ['stationery', 'storage'][i % 2], 'unit_price_minor_units': str(500 + i * 25)} for i in range(1, 31)],
    ]


def _table(name, rows):
    return {'id': name, 'name': name, 'columns': list(rows[0]), 'rows': rows,
            'source': {'kind': 'original_synthetic', 'file': f'{name}.csv', 'truncated': False}}


def build_projects():
    """Return domain -> project, baseline_project, draft_project, expected, mutations, metadata.

    project confirms seven supported rules and retains one unsupported draft.
    expected is a list of {rule_id, table, row, status}; row is one-based or None.
    Mutations record original/new values and explicit affected-rule labels.
    """
    bundles = {}
    for domain, spec in SPECS.items():
        project = new_project(f'Original synthetic {domain} demonstration')
        refs = _references(domain, spec)
        primary = []
        for i in range(200):
            row = {spec['id']: f'{domain[:3].upper()}-{i + 1:04}',
                   spec['number']: str(spec['bounds'][i % 2] if i % 20 < 2 else spec['bounds'][0] + i % 50),
                   spec['enum']: spec['values'][i % 3], spec['dates'][0]: '2026-01-10',
                   spec['dates'][1]: '2026-01-10' if i % 20 == 0 else '2026-01-12',
                   spec['optional']: str([0, 100, 20][i % 3])}
            for (table, column), rows in zip(spec['refs'], refs):
                row[column] = rows[i % len(rows)][column]
            if domain == 'manufacturing':
                row.update(lot_code=f'LOT-{i // 10 + 1:03}', site='Synthetic plant A')
            elif domain == 'healthcare':
                row.update(encounter_class='ambulatory', discharge_summary_id=f'DOC-{i + 1:04}')
            elif domain == 'finance':
                row.update(currency='USD', settlement_reference=f'SET-{i + 1:04}', payment_rail='synthetic_transfer')
            else:
                row.update(tracking_reference=f'TRK-{i + 1:04}', fulfillment_center='Synthetic warehouse A')
            primary.append(row)
        assert len(primary) >= 100
        project['tables'] = [_table(spec['primary'], primary)] + [_table(name, rows) for (name, _), rows in zip(spec['refs'], refs)]
        declarations = [
            ('required', spec['id'], {}, f"{spec['primary']}.{spec['id']} is required."),
            ('range', spec['number'], dict(zip(['min', 'max'], spec['bounds'])), f"{spec['primary']}.{spec['number']} must be between {spec['bounds'][0]} and {spec['bounds'][1]}, inclusive."),
            ('allowed', spec['enum'], {'values': spec['values']}, f"{spec['primary']}.{spec['enum']} must be one of: {', '.join(spec['values'])}."),
            ('date_order', spec['dates'][0], {'other_column': spec['dates'][1]}, f"{spec['primary']}.{spec['dates'][0]} must be on or before {spec['dates'][1]}; equality is allowed."),
        ]
        for table, column in spec['refs']:
            declarations.append(('reference', column, {'target_table': table, 'target_column': column}, f"{spec['primary']}.{column} must reference {table}.{column}."))
        declarations.extend([
            ('range', spec['optional'], {'min': 0, 'max': 100}, f"When supplied, {spec['primary']}.{spec['optional']} must be between 0 and 100, inclusive. A missing value is unknown, not a failure or pass."),
            ('conditional', spec['enum'], {}, spec['conditional']),
        ])
        # Each quote occupies exactly the stated line in policy.md.
        policy_lines = ['# Original synthetic demonstration policy', '',
                        'These authored bounds are demonstration contracts, not clinical thresholds, compliance rules, or business approvals.',
                        'Check only this complete synthetic snapshot. Numeric/date parse failures remain unknown.', '']
        for index, (kind, column, params, text) in enumerate(declarations, 1):
            policy_lines.append(text)
            project['rules'].append({'id': f'{domain}-r{index}', 'text': text, 'table': spec['primary'],
                                     'column': column, 'kind': kind, 'params': params,
                                     'status': 'draft' if kind == 'conditional' else 'confirmed',
                                     'source': {'file': 'policy.md', 'line': len(policy_lines), 'quote': text}})
        project['documents'] = [{'name': 'policy.md', 'segments': [{'line': i, 'text': line} for i, line in enumerate(policy_lines, 1) if line]}]
        baseline = copy.deepcopy(project)
        # Disjoint row mutations: expected labels are declared here, never inferred by the engine.
        catalog = [
            (1, spec['id'], None, 1, 'violation', 'missing_required'),
            (2, spec['number'], str(spec['bounds'][1] + 1), 2, 'violation', 'out_of_range'),
            (3, spec['enum'], 'unrecognized', 3, 'violation', 'invalid_enum'),
            (4, spec['dates'][1], '2026-01-09', 4, 'violation', 'reversed_dates'),
            (5, spec['refs'][0][1], 'ORPHAN-REF-1', 5, 'violation', 'orphan_first_reference'),
            (6, spec['number'], 'not-a-number', 2, 'unknown', 'malformed_numeric'),
            (7, spec['dates'][0], '2026-02-30', 4, 'unknown', 'malformed_date'),
            (8, spec['optional'], None, 7, 'unknown', 'optional_missing'),
            (9, spec['refs'][1][1], 'ORPHAN-REF-2', 6, 'violation', 'orphan_second_reference'),
            (10, spec['dates'][1], None, 4, 'unknown', 'comparison_date_missing'),
        ]
        mutations = []
        overrides = {}
        expanded_catalog = [(row + repeat * 10, column, value, rule, status, label)
                            for repeat in range(5)
                            for row, column, value, rule, status, label in catalog]
        for row_number, column, value, rule_number, status, label in expanded_catalog:
            rule_id = f'{domain}-r{rule_number}'
            row = primary[row_number - 1]
            mutations.append({'row': row_number, 'table': spec['primary'], 'column': column,
                              'before': row[column], 'after': value, 'rule_id': rule_id,
                              'expected_status': status, 'label': label})
            row[column] = value
            overrides[(rule_id, row_number)] = status
        expected = [{'rule_id': rule['id'], 'table': rule['table'], 'row': i,
                     'status': overrides.get((rule['id'], i), 'pass')}
                    for rule in project['rules'][:-1] for i in range(1, 201)]
        expected.append({'rule_id': project['rules'][-1]['id'], 'table': spec['primary'], 'row': None, 'status': 'unknown'})
        draft = copy.deepcopy(project)
        for rule in draft['rules']:
            rule['status'] = 'draft'
        validate_project(project)
        bundles[domain] = {'project': project, 'baseline_project': baseline, 'draft_project': draft,
                           'policy': '\n'.join(policy_lines) + '\n', 'expected': expected, 'mutations': mutations,
                           'metadata': {'origin': 'original_synthetic', 'primary_rows': 200,
                                        'inspiration_urls': [spec['url']], 'copied_source_records': 0,
                                        'scope': 'Schema inspiration only; no real persons, payments, orders or work orders. Flattened healthcare tables are not FHIR resources or FHIR conformance. Finance is not bank/AML compliance. Ecommerce is an authored order-line-like simplification inspired by a wholesale schema. No external accuracy or domain-transfer claim.'}}
    return bundles


def _xlsx(tables):
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    book.properties.created = book.properties.modified = datetime(2026, 1, 1)
    for table in tables:
        sheet = book.create_sheet(table['name'])
        sheet.append(table['columns'])
        for row in table['rows']:
            sheet.append([row.get(column) for column in table['columns']])
    raw = io.BytesIO()
    book.save(raw)
    # Canonical ZIP metadata avoids build-time differences in equivalent fixtures.
    output = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            payload = source.read(name)
            if name == 'docProps/core.xml':
                payload = re.sub(rb'(<dcterms:modified[^>]*>)[^<]+', rb'\g<1>2026-01-01T00:00:00Z', payload)
            target.writestr(info, payload)
    return output.getvalue()


def write_demos(output):
    """Write new directory only; return in-memory build_projects() bundles."""
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    bundles = build_projects()
    for domain, bundle in bundles.items():
        directory = output / domain
        directory.mkdir()
        project = bundle['project']
        for filename, key in [('draft-project.json', 'draft_project'), ('confirmed-project.json', 'project'), ('baseline-project.json', 'baseline_project')]:
            (directory / filename).write_bytes(export_json(bundle[key], include_data=True))
        (directory / 'policy.md').write_text(bundle['policy'], encoding='utf-8')
        for filename, value in [('expectations.json', {'results': bundle['expected'], 'mutations': bundle['mutations']}), ('metadata.json', bundle['metadata'])]:
            (directory / filename).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
        tables = project['tables']
        for table in tables:
            with (directory / f"{table['name']}.csv").open('w', newline='', encoding='utf-8') as handle:
                writer = csv.DictWriter(handle, fieldnames=table['columns'])
                writer.writeheader()
                writer.writerows(table['rows'])
        (directory / 'data.xlsx').write_bytes(_xlsx(tables))
        with sqlite3.connect(directory / 'data.sqlite') as connection:
            for table in tables:
                columns = ', '.join(f'"{column}" TEXT' for column in table['columns'])
                connection.execute(f'CREATE TABLE "{table["name"]}" ({columns})')
                placeholders = ', '.join('?' for _ in table['columns'])
                connection.executemany(f'INSERT INTO "{table["name"]}" VALUES ({placeholders})',
                                       [[row.get(column) for column in table['columns']] for row in table['rows']])
    return bundles


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True, help='New output directory (must not exist)')
    args = parser.parse_args()
    write_demos(args.output)
    print(args.output.resolve())


if __name__ == '__main__':
    main()
