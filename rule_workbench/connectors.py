"""Bounded local imports and read-only PostgreSQL previews; never retain credentials."""
from __future__ import annotations

import csv
import hashlib
import io
import math
import sqlite3
import tempfile
import uuid
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

MAX_UPLOAD = 20 * 1024 * 1024
MAX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 200
MAX_TABLES = 50
MAX_CELLS = 1_000_000


def _payload(payload: bytes) -> None:
    if not isinstance(payload, bytes) or len(payload) > MAX_UPLOAD:
        raise ValueError("Upload must be bytes and at most 20 MB.")


def _columns(values) -> list[str]:
    columns = [str(v) if v is not None else "" for v in values]
    if not columns or len(columns) > MAX_COLUMNS:
        raise ValueError("Tables must have 1–200 columns.")
    if any(not c.strip() for c in columns) or len(set(columns)) != len(columns):
        raise ValueError("Column names must be nonempty and unique.")
    return columns


def _value(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Nonfinite numeric values are unsupported.")
    if isinstance(value, (str, bool, int, float)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    raise ValueError("Unsupported cell type; use text or scalar values.")


def _table(name, columns, records, source):
    rows = []
    truncated = False
    for record in records:
        if len(rows) == MAX_ROWS:
            truncated = True
            break
        if len(record) != len(columns):
            raise ValueError("A data row has a different column count from the header.")
        rows.append(dict(zip(columns, map(_value, record))))
    return {"id": uuid.uuid4().hex, "name": name, "columns": columns,
            "rows": rows, "source": {**source, "truncated": truncated}}


def _check_bundle(tables):
    if len(tables) > MAX_TABLES or sum(len(t["rows"]) * len(t["columns"]) for t in tables) > MAX_CELLS:
        raise ValueError("Import exceeds 50 tables or 1,000,000 cells; split the file.")
    return tables


def read_file(name: str, payload: bytes) -> list[dict]:
    """Import every sheet/table, retaining source hash and explicit truncation."""
    _payload(payload)
    name = Path(name).name
    suffix = Path(name).suffix.lower()
    source = {"file": name, "sha256": hashlib.sha256(payload).hexdigest()}
    try:
        if suffix == ".csv":
            try:
                text = payload.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = payload.decode("cp949")
            records = csv.reader(io.StringIO(text, newline=""), strict=True)
            columns = _columns(next(records, []))
            return [_table(name, columns, records, {**source, "kind": "csv"})]
        if suffix == ".xlsx":
            return _read_xlsx(name, payload, source)
        if suffix in {".db", ".sqlite", ".sqlite3"}:
            return _read_sqlite(payload, source)
        raise ValueError("Supported data files: CSV, XLSX, DB, SQLite, SQLite3.")
    except ValueError:
        raise
    except Exception:
        raise ValueError("Unable to read this data file; verify its format and contents.") from None


def _read_xlsx(name, payload, source):
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED:
            raise ValueError("Expanded XLSX must be at most 100 MB and 10,000 entries.")
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > MAX_TABLES:
            raise ValueError("Import exceeds 50 sheets.")
        tables = []
        for sheet in workbook.worksheets:
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValueError("Tables must have 1–200 columns.")
            records = sheet.iter_rows(values_only=True)
            header = next(records, None)
            if header is None or all(v is None for v in header):
                raise ValueError("Each sheet must have a nonempty header.")
            columns = _columns(header)
            tables.append(_table(f"{name} / {sheet.title}", columns, records,
                                 {**source, "kind": "xlsx", "sheet": sheet.title}))
            _check_bundle(tables)
        return tables
    finally:
        workbook.close()


def _quote(identifier):
    return '"' + identifier.replace('"', '""') + '"'


def _read_sqlite(payload, source):
    if not payload.startswith(b"SQLite format 3\x00"):
        raise ValueError("This is not a SQLite database snapshot.")
    with tempfile.TemporaryDirectory(prefix="rule-workbench-") as directory:
        path = Path(directory) / "snapshot.db"
        path.write_bytes(payload)
        connection = sqlite3.connect(path.as_uri() + "?mode=ro&immutable=1", uri=True)
        try:
            connection.execute("PRAGMA query_only=ON")
            # Bound work even for malicious views/complex storage; views are not imported.
            steps = 0

            def progress():
                nonlocal steps
                steps += 1
                return int(steps > 10000)

            connection.set_progress_handler(progress, 1000)
            names = connection.execute(
                "SELECT name, sql FROM sqlite_master WHERE type='table' "
                "AND name NOT LIKE 'sqlite_%' ORDER BY name LIMIT ?", (MAX_TABLES + 1,)
            ).fetchall()
            if len(names) > MAX_TABLES:
                raise ValueError("Import exceeds 50 tables.")
            tables = []
            for table, definition in names:
                if definition and "VIRTUAL TABLE" in definition.upper():
                    raise ValueError("SQLite virtual tables are unsupported; upload ordinary tables.")
                cursor = connection.execute(f"SELECT * FROM {_quote(table)} LIMIT {MAX_ROWS + 1}")
                columns = _columns([c[0] for c in cursor.description])
                tables.append(_table(table, columns, cursor,
                                     {**source, "kind": "sqlite", "table": table}))
                _check_bundle(tables)
            return tables
        finally:
            connection.close()


def _engine(url):
    from sqlalchemy import create_engine
    from sqlalchemy.engine import make_url

    try:
        parsed = make_url(url)
        if parsed.drivername not in {"postgresql", "postgresql+psycopg"}:
            raise ValueError
        # Arbitrary libpq query options (service/passfile/options) are not accepted.
        if parsed.query:
            raise ValueError
        return create_engine(parsed.set(drivername="postgresql+psycopg"),
                             connect_args={"connect_timeout": 5}, pool_pre_ping=False)
    except Exception:
        raise ValueError("Use a PostgreSQL URL without query options.") from None


def _readonly(connection):
    from sqlalchemy import text

    connection.execute(text("SET TRANSACTION READ ONLY"))
    connection.execute(text("SET LOCAL statement_timeout = '5000ms'"))
    connection.execute(text("SET LOCAL lock_timeout = '2000ms'"))


def list_db_tables(url: str) -> list[str]:
    """List ordinary tables in the connection's default schema, with safe errors."""
    from sqlalchemy import inspect

    engine = _engine(url)
    try:
        with engine.connect() as connection, connection.begin():
            _readonly(connection)
            names = inspect(connection).get_table_names()
            if len(names) > MAX_TABLES:
                raise ValueError
            return sorted(names)
    except Exception:
        raise ValueError("Unable to list PostgreSQL tables (maximum 50); verify connection and permissions.") from None
    finally:
        engine.dispose()


def read_db_table(url: str, table: str) -> dict:
    """Read at most 5,000 rows using a validated, dialect-quoted identifier."""
    from sqlalchemy import inspect, text

    engine = _engine(url)
    try:
        with engine.connect() as connection, connection.begin():
            _readonly(connection)
            if table not in inspect(connection).get_table_names():
                raise ValueError
            quoted = connection.dialect.identifier_preparer.quote_identifier(table)
            result = connection.execution_options(stream_results=True).execute(
                text(f"SELECT * FROM {quoted} LIMIT {MAX_ROWS + 1}")
            )
            try:
                columns = _columns(result.keys())
                return _table(table, columns, result, {"kind": "postgresql", "table": table})
            finally:
                result.close()
    except Exception:
        raise ValueError("Unable to read PostgreSQL table; verify name, permissions and scalar column types.") from None
    finally:
        engine.dispose()
